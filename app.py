"""Buscador de salões de beleza no Rio de Janeiro (interface web em Streamlit).

Rodar localmente:  streamlit run app.py
Configuração (em .streamlit/secrets.toml ou nos Secrets do Streamlit Cloud):
  GOOGLE_MAPS_API_KEY = "..."      obrigatório
  MAX_BUSCAS_DIA = 15              opcional: limite de buscas novas por dia, somando todos os usuários
  APP_SENHA = "..."                opcional: exige senha para usar o app
"""
import base64
import io
import os
import re
from datetime import date

import pandas as pd
import pydeck as pdk
import streamlit as st

import buscar_saloes as core

PASTA = os.path.dirname(os.path.abspath(__file__))

st.set_page_config(page_title="Buscador Vitallis", page_icon=os.path.join(PASTA, "logo_vitallis_profissional.png"), layout="wide")


def config(nome, padrao=None):
    try:
        if nome in st.secrets:
            return st.secrets[nome]
    except Exception:
        pass
    return os.environ.get(nome, padrao)


CHAVE = config("GOOGLE_MAPS_API_KEY")
MAX_BUSCAS_DIA = int(config("MAX_BUSCAS_DIA", 15))
SENHA = config("APP_SENHA")

TERMOS = {
    "Salão de beleza": "salão de beleza",
    "Cabeleireiro": "cabeleireiro",
    "Manicure e pedicure": "manicure pedicure esmalteria",
    "Barbearia": "barbearia",
    "Estética e sobrancelha": "estética sobrancelha cílios",
}


# ---------------------------------------------------------------- limite de uso

@st.cache_resource
def contador():
    # Compartilhado entre todos os usuários enquanto o servidor estiver no ar.
    return {"dia": date.today(), "buscas": 0}


def pode_buscar():
    c = contador()
    if c["dia"] != date.today():
        c["dia"], c["buscas"] = date.today(), 0
    return c["buscas"] < MAX_BUSCAS_DIA


def registrar_busca():
    contador()["buscas"] += 1


# ---------------------------------------------------------------- busca (com cache)

class ErroBusca(Exception):
    pass


def normalizar_local(texto):
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) == 8 and len(texto.strip()) <= 10:
        return f"CEP {digitos[:5]}-{digitos[5:]}"
    return texto.strip()


@st.cache_data(ttl=6 * 3600, show_spinner=False)
def geocodificar(local):
    try:
        return core.geocodificar(local, "google", CHAVE)
    except SystemExit as e:
        raise ErroBusca(str(e))


@st.cache_data(ttl=6 * 3600, show_spinner=False)
def buscar(lat, lon, raio, termo):
    try:
        resultados = core.buscar_google(lat, lon, raio, CHAVE, termo)
    except SystemExit as e:
        raise ErroBusca(str(e))
    vistos, unicos = set(), []
    for r in resultados:
        if r["latitude"] is None:
            continue
        r["distancia_m"] = core.distancia_m(lat, lon, r["latitude"], r["longitude"])
        chave_dup = (core.sem_acento(r["nome"]), round(r["latitude"], 4), round(r["longitude"], 4))
        if r["distancia_m"] <= raio and chave_dup not in vistos:
            vistos.add(chave_dup)
            unicos.append(r)
    unicos.sort(key=lambda r: r["distancia_m"])
    return unicos


def gerar_xlsx(df, titulo):
    from openpyxl.styles import Font, PatternFill

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, startrow=2, sheet_name="Salões")
        ws = w.sheets["Salões"]
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=13)
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="8E3B6B")
        larguras = {"Nome": 34, "Endereço": 48, "Horário": 50, "Site": 34, "Mapa": 40, "Tipo": 20, "Telefone": 18}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[ws.cell(3, i).column_letter].width = larguras.get(col, 12)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{ws.cell(3, len(df.columns)).column_letter}{ws.max_row}"
    return buf.getvalue()


# ---------------------------------------------------------------- interface

def logo_b64(nome):
    with open(os.path.join(PASTA, nome), "rb") as f:
        return base64.b64encode(f.read()).decode()


st.markdown(f"""
<style>
  .block-container {{ padding-top: 1.5rem; }}
  header[data-testid="stHeader"] {{ display: none; }}
  .vt-banner {{
    display: flex; align-items: center; gap: 28px; padding: 22px 32px; margin-bottom: 18px;
    border-radius: 16px; background: radial-gradient(circle at 15% 50%, #3a2c14 0%, #17130c 55%, #0e0c08 100%);
    border: 1px solid #6b5426;
  }}
  .vt-banner img {{ height: 104px; }}
  .vt-banner h1 {{
    margin: 0; padding: 0; font-family: Georgia, 'Times New Roman', serif; font-weight: 500; font-size: 2.5rem;
    letter-spacing: .04em; background: linear-gradient(90deg, #f6e3a8, #c89b3c 55%, #f1d58a);
    -webkit-background-clip: text; background-clip: text; color: transparent;
  }}
  .vt-banner p {{ margin: 6px 0 0; color: #d9ccb0; font-size: 1rem; }}
  @media (max-width: 640px) {{
    .vt-banner {{ flex-direction: column; text-align: center; gap: 10px; padding: 18px; }}
    .vt-banner img {{ height: 80px; }}
    .vt-banner h1 {{ font-size: 1.9rem; }}
  }}
</style>
<div class="vt-banner">
  <img src="data:image/png;base64,{logo_b64('logo_vitallis_profissional.png')}" alt="Vitallis Profissional">
  <div>
    <h1>Buscador Vitallis</h1>
    <p>Encontre salões de beleza perto de você no Rio de Janeiro. É só digitar o CEP.</p>
  </div>
</div>
""", unsafe_allow_html=True)

if not CHAVE:
    st.error("O app não está configurado: falta a chave GOOGLE_MAPS_API_KEY.")
    st.stop()

if SENHA and not st.session_state.get("liberado"):
    s = st.text_input("Senha de acesso", type="password")
    if s:
        if s == SENHA:
            st.session_state.liberado = True
            st.rerun()
        st.error("Senha incorreta.")
    st.stop()

with st.form("busca"):
    c1, c2, c3 = st.columns([3, 1.4, 1.6])
    local = c1.text_input("CEP", placeholder="Ex.: 20510-130", help="Também aceita endereço ou nome do bairro.")
    raio_km = c2.select_slider("Raio", options=[0.5, 1.0, 1.5, 2.0, 3.0, 5.0], value=1.0,
                               format_func=lambda v: f"{v:g} km")
    tipo = c3.selectbox("Tipo", list(TERMOS))
    enviar = st.form_submit_button("Buscar salões", type="primary", width="stretch")

if enviar:
    if not local.strip():
        st.warning("Digite um CEP.")
        st.stop()
    consulta = normalizar_local(local)
    raio = int(raio_km * 1000)
    ja_em_cache = st.session_state.get("ultima") == (consulta, raio, tipo)
    if not ja_em_cache and not pode_buscar():
        st.error("O limite diário de buscas do app foi atingido. Tente novamente amanhã.")
        st.stop()
    with st.spinner("Buscando salões..."):
        try:
            lat, lon, origem = geocodificar(consulta)
            resultados = buscar(lat, lon, raio, TERMOS[tipo])
        except ErroBusca as e:
            st.error(str(e))
            st.stop()
    if not ja_em_cache:
        registrar_busca()
    st.session_state.update(ultima=(consulta, raio, tipo), resultado=(lat, lon, origem, raio, tipo, resultados))

if "resultado" not in st.session_state:
    st.info("Digite o CEP e clique em **Buscar salões**.")
    st.stop()

lat, lon, origem, raio, tipo, resultados = st.session_state.resultado
st.subheader(f"{len(resultados)} resultados · {tipo.lower()} a até {raio / 1000:g} km")
st.caption(f"📍 {origem}")

if not resultados:
    st.warning("Nenhum salão encontrado nesse raio. Tente aumentar a distância.")
    st.stop()

df = pd.DataFrame(resultados)
df["nota"] = pd.to_numeric(df["nota"], errors="coerce")
df["avaliacoes"] = pd.to_numeric(df["avaliacoes"], errors="coerce").fillna(0).astype(int)

f1, f2 = st.columns([1, 1])
nota_min = f1.slider("Nota mínima", 0.0, 5.0, 0.0, 0.5, format="%.1f")
ordem = f2.radio("Ordenar por", ["Mais perto", "Melhor nota", "Mais avaliações"], horizontal=True)
df = df[df["nota"].fillna(0) >= nota_min]
if ordem == "Melhor nota":
    df = df.sort_values(["nota", "avaliacoes"], ascending=False)
elif ordem == "Mais avaliações":
    df = df.sort_values("avaliacoes", ascending=False)

aba_lista, aba_mapa = st.columns([1.35, 1], gap="medium")

df = df.reset_index(drop=True)
df["n"] = df.index + 1
tabela = pd.DataFrame({
    "Nº": df["n"],
    "Nome": df["nome"],
    "Nota": df["nota"],
    "Avaliações": df["avaliacoes"],
    "Distância (m)": df["distancia_m"],
    "Endereço": df["endereco"],
    "Telefone": df["telefone"],
    "Horário": df["horario"],
    "Site": df["site"].replace("", None),
    "Mapa": df["link_mapa"].replace("", None),
    "Tipo": df["tipo"],
})

with aba_lista:
    st.caption("Clique em uma linha para destacar o salão no mapa.")
    evento = st.dataframe(
        tabela, hide_index=True, width="stretch", height=560,
        on_select="rerun", selection_mode="single-row", key="grade",
        column_config={
            "Nº": st.column_config.NumberColumn(width="small"),
            "Nota": st.column_config.NumberColumn(format="%.1f ⭐"),
            "Site": st.column_config.LinkColumn(display_text="abrir"),
            "Mapa": st.column_config.LinkColumn(display_text="Google Maps"),
            "Horário": st.column_config.TextColumn(width="large"),
        },
    )
    linhas = evento.selection.rows if evento and evento.selection else []
    selecionado = int(linhas[0]) if linhas else None

with aba_mapa:
    pontos = df.assign(
        nota_txt=df["nota"].map(lambda n: f"{n:.1f} ⭐" if pd.notna(n) else "sem nota"),
        rotulo=df["n"].astype(str),
    )
    centro_lat, centro_lon = lat, lon
    zoom = 15 if raio <= 1000 else 14 if raio <= 2000 else 13
    if selecionado is not None:
        centro_lat, centro_lon = df.loc[selecionado, "latitude"], df.loc[selecionado, "longitude"]
        zoom = 16.5
    camadas = [
        pdk.Layer("ScatterplotLayer", data=pontos, get_position="[longitude, latitude]", get_radius=14,
                  radius_min_pixels=9, get_fill_color=[184, 134, 43, 235], get_line_color=[40, 30, 10],
                  stroked=True, line_width_min_pixels=1, pickable=True),
        pdk.Layer("TextLayer", data=pontos, get_position="[longitude, latitude]", get_text="rotulo",
                  get_size=11, get_color=[255, 255, 255], get_alignment_baseline="'center'",
                  font_weight=700),
        pdk.Layer("ScatterplotLayer", data=pd.DataFrame([{"latitude": lat, "longitude": lon}]),
                  get_position="[longitude, latitude]", get_radius=20, radius_min_pixels=8,
                  get_fill_color=[30, 110, 220, 255], get_line_color=[255, 255, 255], stroked=True,
                  line_width_min_pixels=2),
    ]
    if selecionado is not None:
        camadas.append(pdk.Layer(
            "ScatterplotLayer", data=pontos.iloc[[selecionado]], get_position="[longitude, latitude]",
            get_radius=30, radius_min_pixels=16, filled=False, stroked=True,
            get_line_color=[220, 40, 90], line_width_min_pixels=3))
    st.pydeck_chart(pdk.Deck(
        layers=camadas, map_style="light",
        initial_view_state=pdk.ViewState(latitude=centro_lat, longitude=centro_lon, zoom=zoom),
        tooltip={"html": "<b>{n}. {nome}</b><br/>{nota_txt} · {distancia_m} m<br/>{endereco}<br/>{telefone}"},
    ), height=560)
    st.caption("🔵 CEP buscado · 🟡 salões (o número é o mesmo da grade)")
    if selecionado is not None:
        r = df.loc[selecionado]
        st.markdown(f"**{r['n']}. {r['nome']}** · {r['endereco']}  \n"
                    f"📞 {r['telefone'] or 'sem telefone'} · [Abrir no Google Maps]({r['link_mapa']})")

titulo = f"Salões perto de {origem} · raio {raio / 1000:g} km · {tipo}"
nome_arq = f"saloes_{core.sem_acento(origem.split(',')[0]).replace(' ', '_')[:30]}"
d1, d2, _ = st.columns([1, 1, 3])
d1.download_button("⬇️ Baixar Excel", gerar_xlsx(tabela, titulo), f"{nome_arq}.xlsx",
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
d2.download_button("⬇️ Baixar CSV", tabela.to_csv(index=False).encode("utf-8-sig"), f"{nome_arq}.csv",
                   "text/csv", width="stretch")

st.divider()
st.caption("Dados: Google Maps. Informações como horário e telefone podem estar desatualizadas; confirme com o salão.")
