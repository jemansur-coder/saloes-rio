#!/usr/bin/env python3
"""Localizador de salões de beleza no município do Rio de Janeiro.

Busca salões perto de um bairro/endereço (ou coordenadas) dentro de um raio
e gera um relatório em CSV e XLSX, ordenado por distância.

Fontes:
  google  Google Places API (New). Precisa de chave em GOOGLE_MAPS_API_KEY
          (ou --chave). Traz nota, nº de avaliações e link do Maps.
  osm     OpenStreetMap (Nominatim + Overpass). Gratuito, sem chave,
          mas com cobertura menor e sem avaliações.
  auto    (padrão) usa google se houver chave, senão osm.

Exemplos:
  python3 buscar_saloes.py --local Centro
  python3 buscar_saloes.py --local "Rua Barata Ribeiro 200, Copacabana" --raio 800
  python3 buscar_saloes.py --lat -22.9068 --lon -43.1729 --raio 1500 --fonte osm
  GOOGLE_MAPS_API_KEY=... python3 buscar_saloes.py --local Tijuca --saida relatorios/

Dependências: requests, openpyxl (pip install requests openpyxl)
"""
import argparse
import csv
import math
import os
import sys
import time
from datetime import datetime

import requests

USER_AGENT = "buscar-saloes-rio/1.0 (uso pessoal)"
TIMEOUT = 60

# Centro aproximado de alguns bairros, usado quando o geocodificador não responde.
BAIRROS = {
    "centro": (-22.9068, -43.1729),
    "lapa": (-22.9133, -43.1800),
    "flamengo": (-22.9329, -43.1760),
    "laranjeiras": (-22.9350, -43.1880),
    "botafogo": (-22.9519, -43.1857),
    "copacabana": (-22.9711, -43.1822),
    "ipanema": (-22.9838, -43.2096),
    "leblon": (-22.9847, -43.2233),
    "tijuca": (-22.9249, -43.2326),
    "vila isabel": (-22.9160, -43.2440),
    "grajau": (-22.9220, -43.2610),
    "meier": (-22.9020, -43.2800),
    "madureira": (-22.8726, -43.3375),
    "jacarepagua": (-22.9656, -43.3930),
    "barra da tijuca": (-23.0004, -43.3659),
    "recreio dos bandeirantes": (-23.0180, -43.4640),
    "campo grande": (-22.9035, -43.5591),
}

COLUNAS = [
    "nome", "tipo", "endereco", "bairro", "telefone", "horario", "site",
    "nota", "avaliacoes", "distancia_m", "latitude", "longitude", "link_mapa", "fonte",
]


def sem_acento(s):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").lower().strip()


def distancia_m(lat1, lon1, lat2, lon2):
    r = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return round(2 * r * math.asin(math.sqrt(a)))


# ---------------------------------------------------------------- geocodificação

def geocodificar(local, fonte, chave):
    consulta = f"{local}, Rio de Janeiro, RJ, Brasil"
    try:
        if fonte == "google":
            r = requests.post(
                "https://places.googleapis.com/v1/places:searchText",
                headers={"X-Goog-Api-Key": chave, "X-Goog-FieldMask": "places.location,places.formattedAddress"},
                json={"textQuery": consulta, "languageCode": "pt-BR", "regionCode": "BR"},
                timeout=TIMEOUT,
            )
            r.raise_for_status()
            lugares = r.json().get("places", [])
            if lugares:
                loc = lugares[0]["location"]
                return loc["latitude"], loc["longitude"], lugares[0].get("formattedAddress", consulta)
        else:
            r = requests.get(
                "https://nominatim.openstreetmap.org/search",
                params={"q": consulta, "format": "json", "limit": 1, "countrycodes": "br"},
                headers={"User-Agent": USER_AGENT},
                timeout=TIMEOUT,
            )
            r.raise_for_status()
            res = r.json()
            if res:
                return float(res[0]["lat"]), float(res[0]["lon"]), res[0].get("display_name", consulta)
    except requests.RequestException as e:
        print(f"Aviso: geocodificação falhou ({e}).", file=sys.stderr)
    chave_bairro = sem_acento(local)
    if chave_bairro in BAIRROS:
        lat, lon = BAIRROS[chave_bairro]
        return lat, lon, f"{local} (centro aproximado do bairro)"
    sys.exit(f"Não encontrei '{local}'. Tente um endereço mais completo ou use --lat/--lon.")


# ---------------------------------------------------------------- Google Places

GOOGLE_CAMPOS = ",".join(f"places.{c}" for c in [
    "displayName", "primaryTypeDisplayName", "formattedAddress", "addressComponents",
    "nationalPhoneNumber", "regularOpeningHours.weekdayDescriptions", "websiteUri",
    "rating", "userRatingCount", "location", "googleMapsUri", "businessStatus",
]) + ",nextPageToken"


def buscar_google(lat, lon, raio, chave, termo):
    resultados, token = [], None
    for _ in range(3):  # a API devolve no máximo 60 resultados (3 páginas de 20)
        corpo = {
            "textQuery": termo,
            "languageCode": "pt-BR",
            "regionCode": "BR",
            "pageSize": 20,
            "locationBias": {"circle": {"center": {"latitude": lat, "longitude": lon}, "radius": float(raio)}},
        }
        if token:
            corpo["pageToken"] = token
        r = requests.post(
            "https://places.googleapis.com/v1/places:searchText",
            headers={"X-Goog-Api-Key": chave, "X-Goog-FieldMask": GOOGLE_CAMPOS},
            json=corpo, timeout=TIMEOUT,
        )
        if r.status_code != 200:
            sys.exit(f"Erro da Google Places API ({r.status_code}): {r.text[:300]}")
        dados = r.json()
        for p in dados.get("places", []):
            if p.get("businessStatus") == "CLOSED_PERMANENTLY":
                continue
            loc = p.get("location", {})
            bairro = next((c.get("longText") for c in p.get("addressComponents", [])
                           if "sublocality" in c.get("types", []) or "sublocality_level_1" in c.get("types", [])), "")
            resultados.append({
                "nome": p.get("displayName", {}).get("text", ""),
                "tipo": p.get("primaryTypeDisplayName", {}).get("text", ""),
                "endereco": p.get("formattedAddress", ""),
                "bairro": bairro,
                "telefone": p.get("nationalPhoneNumber", ""),
                "horario": " | ".join(p.get("regularOpeningHours", {}).get("weekdayDescriptions", [])),
                "site": p.get("websiteUri", ""),
                "nota": p.get("rating", ""),
                "avaliacoes": p.get("userRatingCount", ""),
                "latitude": loc.get("latitude"),
                "longitude": loc.get("longitude"),
                "link_mapa": p.get("googleMapsUri", ""),
                "fonte": "Google Places",
            })
        token = dados.get("nextPageToken")
        if not token:
            break
        time.sleep(2)
    return resultados


# ---------------------------------------------------------------- OpenStreetMap

OVERPASS_URLS = [
    "https://overpass-api.de/api/interpreter",
    "https://overpass.kumi.systems/api/interpreter",
]
TIPOS_OSM = {"hairdresser": "Cabeleireiro", "beauty": "Salão de beleza", "cosmetics": "Cosméticos",
             "nails": "Manicure", "barber": "Barbearia"}


def buscar_osm(lat, lon, raio):
    consulta = f"""
[out:json][timeout:60];
(
  nwr(around:{raio},{lat},{lon})["shop"~"^(hairdresser|beauty)$"];
  nwr(around:{raio},{lat},{lon})["amenity"="beauty_salon"];
);
out center tags;
"""
    ultimo_erro = None
    for url in OVERPASS_URLS:
        try:
            r = requests.post(url, data={"data": consulta}, headers={"User-Agent": USER_AGENT}, timeout=TIMEOUT + 30)
            r.raise_for_status()
            elementos = r.json().get("elements", [])
            break
        except requests.RequestException as e:
            ultimo_erro = e
    else:
        sys.exit(f"Não consegui acessar o Overpass (OpenStreetMap): {ultimo_erro}")

    resultados = []
    for el in elementos:
        t = el.get("tags", {})
        la = el.get("lat") or el.get("center", {}).get("lat")
        lo = el.get("lon") or el.get("center", {}).get("lon")
        rua = " ".join(x for x in [t.get("addr:street", ""), t.get("addr:housenumber", "")] if x)
        endereco = ", ".join(x for x in [rua, t.get("addr:suburb", ""), t.get("addr:postcode", "")] if x)
        tipo = TIPOS_OSM.get(t.get("shop", ""), "Salão de beleza")
        if t.get("beauty"):
            tipo += f" ({t['beauty']})"
        resultados.append({
            "nome": t.get("name", "(sem nome no mapa)"),
            "tipo": tipo,
            "endereco": endereco,
            "bairro": t.get("addr:suburb", ""),
            "telefone": t.get("phone") or t.get("contact:phone") or t.get("contact:whatsapp", ""),
            "horario": t.get("opening_hours", ""),
            "site": t.get("website") or t.get("contact:website") or t.get("contact:instagram", ""),
            "nota": "",
            "avaliacoes": "",
            "latitude": la,
            "longitude": lo,
            "link_mapa": f"https://www.openstreetmap.org/{el['type']}/{el['id']}",
            "fonte": "OpenStreetMap",
        })
    return resultados


# ---------------------------------------------------------------- relatório

def salvar(resultados, prefixo, origem_txt, raio, fonte):
    os.makedirs(os.path.dirname(prefixo) or ".", exist_ok=True)
    with open(prefixo + ".csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUNAS, extrasaction="ignore")
        w.writeheader()
        w.writerows(resultados)
    arquivos = [prefixo + ".csv"]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Aviso: openpyxl não instalado, gerei só o CSV.", file=sys.stderr)
        return arquivos

    wb = Workbook()
    ws = wb.active
    ws.title = "Salões"
    ws.append([f"Salões de beleza perto de: {origem_txt}"])
    ws.append([f"Raio: {raio} m · Fonte: {fonte} · Gerado em {datetime.now():%d/%m/%Y %H:%M} · {len(resultados)} resultados"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(COLUNAS)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8E3B6B")
    for r in resultados:
        ws.append([r.get(c, "") for c in COLUNAS])
    larguras = {"nome": 34, "tipo": 20, "endereco": 48, "bairro": 16, "telefone": 18, "horario": 50,
                "site": 34, "link_mapa": 40}
    for i, col in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(col, 12)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(COLUNAS))}{ws.max_row}"
    wb.save(prefixo + ".xlsx")
    arquivos.append(prefixo + ".xlsx")
    return arquivos


def main():
    ap = argparse.ArgumentParser(description="Busca salões de beleza no Rio de Janeiro e gera relatório.")
    ap.add_argument("--local", default="Centro", help="Bairro ou endereço (padrão: Centro)")
    ap.add_argument("--lat", type=float, help="Latitude (ignora --local)")
    ap.add_argument("--lon", type=float, help="Longitude (ignora --local)")
    ap.add_argument("--raio", type=int, default=1500, help="Raio em metros (padrão: 1500)")
    ap.add_argument("--fonte", choices=["auto", "google", "osm"], default="auto")
    ap.add_argument("--chave", default=os.environ.get("GOOGLE_MAPS_API_KEY"), help="Chave da Google Places API")
    ap.add_argument("--termo", default="salão de beleza", help="Termo de busca no Google (padrão: salão de beleza)")
    ap.add_argument("--limite", type=int, default=0, help="Mostrar só os N mais próximos (0 = todos)")
    ap.add_argument("--saida", default=".", help="Pasta onde salvar o relatório")
    a = ap.parse_args()

    fonte = a.fonte if a.fonte != "auto" else ("google" if a.chave else "osm")
    if fonte == "google" and not a.chave:
        sys.exit("Fonte google exige chave: defina GOOGLE_MAPS_API_KEY ou use --chave.")

    if a.lat is not None and a.lon is not None:
        lat, lon, origem = a.lat, a.lon, f"{a.lat}, {a.lon}"
    else:
        lat, lon, origem = geocodificar(a.local, fonte, a.chave)
    print(f"Origem: {origem} ({lat:.5f}, {lon:.5f}) · raio {a.raio} m · fonte {fonte}")

    resultados = buscar_google(lat, lon, a.raio, a.chave, a.termo) if fonte == "google" else buscar_osm(lat, lon, a.raio)

    vistos, unicos = set(), []
    for r in resultados:
        if r["latitude"] is None:
            continue
        r["distancia_m"] = distancia_m(lat, lon, r["latitude"], r["longitude"])
        chave_dup = (sem_acento(r["nome"]), round(r["latitude"], 4), round(r["longitude"], 4))
        if r["distancia_m"] <= a.raio and chave_dup not in vistos:
            vistos.add(chave_dup)
            unicos.append(r)
    unicos.sort(key=lambda r: r["distancia_m"])
    if a.limite:
        unicos = unicos[: a.limite]

    nome_local = sem_acento(a.local if a.lat is None else "coordenadas").replace(" ", "_").replace(",", "")[:40]
    prefixo = os.path.join(a.saida, f"saloes_{nome_local}_{a.raio}m_{fonte}_{datetime.now():%Y%m%d_%H%M%S}")
    arquivos = salvar(unicos, prefixo, origem, a.raio, fonte)

    print(f"\n{len(unicos)} salões encontrados. Os mais próximos:")
    for r in unicos[:10]:
        extra = f" · nota {r['nota']} ({r['avaliacoes']})" if r["nota"] != "" else ""
        print(f"  {r['distancia_m']:>5} m  {r['nome']} · {r['endereco'] or 'endereço não informado'}"
              f"{' · ' + r['telefone'] if r['telefone'] else ''}{extra}")
    print("\nRelatório salvo em:\n  " + "\n  ".join(arquivos))


if __name__ == "__main__":
    main()
